"""Machine document checks and all-page rendering, distinct from visual review.

Usage: python inspect_documents.py --render-dir /scratch/manuscript-pages
Requires pypdf, pypdfium2 and python-docx. Does not calculate scientific results.
"""
import argparse
import hashlib
import json
import re
from pathlib import Path
from docx import Document
from pypdf import PdfReader
import pypdfium2 as pdfium

BASE=Path(__file__).resolve().parents[1]
NAMES=['manuscript_english','manuscript_chinese','supporting_information_english','supporting_information_chinese']

def inspect(render_dir):
    records=[]
    for name in NAMES:
        docx=BASE/'documents'/(name+'.docx');pdf=docx.with_suffix('.pdf')
        doc=Document(docx);reader=PdfReader(pdf);images=pdfium.PdfDocument(pdf)
        out=render_dir/name;out.mkdir(parents=True,exist_ok=True)
        pages=[];alltext=[]
        for i,page in enumerate(reader.pages):
            text=page.extract_text();alltext.append(text)
            p=images[i];bitmap=p.render(scale=1.5);im=bitmap.to_pil()
            target=out/f'page-{i+1:03}.png';im.save(target)
            pages.append({'page':i+1,'text_characters':len(text),'rendered_pixels':[im.width,im.height],
                          'image_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),
                          'letter_page':abs(float(page.mediabox.width)-612)<1 and abs(float(page.mediabox.height)-792)<1})
            bitmap.close();p.close()
        images.close()
        joined='\n'.join(alltext)
        xml=doc._element.xml
        expected_figures=6 if name.startswith('manuscript') else 0
        checks={
            'all_pages_letter':all(p['letter_page'] for p in pages),
            'no_empty_pages':all(p['text_characters']>25 for p in pages),
            'no_replacement_character':'\ufffd' not in joined,
            'no_unresolved_references':'[@' not in joined and 'VERIFIED_REFERENCES' not in joined,
            'no_raw_display_latex':not any(x in joined for x in ['\\widehat','\\begin{','$$','\\langle']),
            'author_present':('Siyi Song' in joined or '宋思毅' in joined),
            'all_expected_main_figure_captions': all(re.search(r'(?:Figure|图)\s*'+str(i),joined) for i in range(1,expected_figures+1)),
            'document_has_hyperlinks':'w:hyperlink' in xml,
            'all_pages_rendered':len(list(out.glob('page-*.png')))==len(pages),
        }
        records.append({'name':name,'pdf':pdf.relative_to(BASE.parent).as_posix(),'docx':docx.relative_to(BASE.parent).as_posix(),
                        'pdf_sha256':hashlib.sha256(pdf.read_bytes()).hexdigest(),'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),
                        'page_count':len(pages),'paragraph_count':len(doc.paragraphs),'table_count':len(doc.tables),
                        'inline_shapes':len(doc.inline_shapes),'pdf_text_characters':len(joined),'pages':pages,'checks':checks})
        (out/'extracted.txt').write_text(joined,encoding='utf-8')
    result={'scope':'Machine structure/text checks and all-page PNG rendering. Visual review is recorded separately; neither substitutes for scientific or human peer review.',
            'passed':all(all(r['checks'].values()) for r in records),'documents':records}
    (BASE/'results'/'document_qa.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'passed':result['passed'],'documents':[{k:r[k] for k in ['name','page_count','checks']} for r in records]},ensure_ascii=False,indent=2))
    return 0 if result['passed'] else 1

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--render-dir',type=Path,required=True);args=ap.parse_args()
    raise SystemExit(inspect(args.render_dir))
