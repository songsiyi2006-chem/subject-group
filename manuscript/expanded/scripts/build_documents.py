"""Build two expanded journal-style DOCX editions from reference-resolved Markdown.

Requires python-docx/Pillow and Matplotlib for equation assets.
Run --prepare-equations first, render them with render_equations.py, then run
without arguments. PDF export is separately performed by Word or LibreOffice.
No scientific calculation is performed by this script.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import re
import textwrap
from pathlib import Path
from urllib.parse import quote
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[3]
BASE = ROOT / 'manuscript' / 'expanded'
NAMES = ['manuscript_english', 'manuscript_chinese']
EQUATIONS = BASE / 'assets' / 'equations'
DISPLAY = re.compile(r'(?ms)^\$\$\s*\n(.*?)\n\$\$\s*$|^\\\[\s*\n(.*?)\n\\\]\s*$')
SOURCE_COMMIT = '1ad05c243155a9f64b18e0d58c665424fde919a0'

def equation_id(tex):
    return 'eq_' + hashlib.sha256(tex.strip().encode()).hexdigest()[:16]

def prepare():
    result = {}
    for name in NAMES:
        text = (BASE / 'rendered' / (name+'.md')).read_text('utf-8')
        for m in DISPLAY.finditer(text):
            tex = (m.group(1) or m.group(2)).strip()
            result[equation_id(tex)] = tex
    EQUATIONS.mkdir(parents=True,exist_ok=True)
    (EQUATIONS/'equations.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'equations_prepared':len(result)}))

def font(run, size=None, bold=None, mono=False):
    run.font.name = 'Consolas' if mono else 'Times New Roman'
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    run.font.color.rgb = RGBColor(0,0,0)
    rp=run._element.get_or_add_rPr()
    rf=rp.rFonts
    if rf is None:
        rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    rf.set(qn('w:eastAsia'),'SimSun')
    rf.set(qn('w:cs'),'Times New Roman')

def url_target(target, source):
    if re.match(r'https?://|mailto:',target):return target
    if target.startswith('#'):return target
    p=(source.parent/target.split('#')[0]).resolve()
    rel=p.relative_to(ROOT).as_posix()
    ref='main' if rel.startswith('manuscript/') else SOURCE_COMMIT
    return f'https://github.com/songsiyi2006-chem/subject-group/blob/{ref}/'+quote(rel,safe='/')

def hyperlink(p,label,url,size=None):
    h=OxmlElement('w:hyperlink')
    h.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
    r=OxmlElement('w:r');prop=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman');fonts.set(qn('w:eastAsia'),'SimSun');prop.append(fonts)
    color=OxmlElement('w:color');color.set(qn('w:val'),'000000');prop.append(color)
    under=OxmlElement('w:u');under.set(qn('w:val'),'single');prop.append(under)
    if size:
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(round(size*2)));prop.append(sz)
    r.append(prop);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def inline(p,text,source,size=None):
    pattern=re.compile(r'(\[[^\]]+\]\([^)]+\)|\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)')
    for part in pattern.split(text):
        if not part:continue
        m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
        if m:hyperlink(p,m.group(1),url_target(m.group(2),source),size);continue
        bold=part.startswith('**') and part.endswith('**')
        mono=part.startswith('`') and part.endswith('`')
        italic=part.startswith('*') and part.endswith('*') and not bold
        if bold:part=part[2:-2]
        elif mono:part=part[1:-1]
        elif italic:part=part[1:-1]
        run=p.add_run(part)
        font(run,size,True if bold else None,mono)
        if italic:run.italic=True

def style_doc(doc, chinese, supporting):
    sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.78)
    sec.left_margin=Inches(.90);sec.right_margin=Inches(.90)
    sec.header_distance=Inches(.32);sec.footer_distance=Inches(.32)
    for key in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        st=doc.styles[key];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'SimSun')
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.15
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    doc.styles['Title'].font.size=Pt(19)
    doc.styles['Title'].font.bold=True
    doc.styles['Title'].paragraph_format.space_after=Pt(12)
    for name,size,before,after in [('Heading 1',13,13,6),('Heading 2',11.5,10,5),('Heading 3',11,8,4)]:
        s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True
        s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after)
        s.paragraph_format.keep_with_next=True
    doc.styles['Caption'].font.size=Pt(9.5)
    doc.styles['Caption'].font.italic=False
    doc.styles['Caption'].font.bold=False
    # The bundled blank template can carry a blue Title paragraph border.
    # Remove paragraph rules from both styles and document content explicitly.
    for element in [doc.styles.element,doc._element]:
        for border in element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    hp=sec.header.paragraphs[0]
    label=('补充信息' if supporting else '计算化学研究论文') if chinese else ('Supporting information' if supporting else 'Computational chemistry research manuscript')
    font(hp.add_run(label),8)
    hp.paragraph_format.space_after=Pt(0)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(fp.add_run('宋思毅 · ' if chinese else 'Siyi Song · '),8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
    doc.core_properties.author='Siyi Song / 宋思毅'
    doc.core_properties.subject='Reference and observable dependence in molecular learning and multiscale computational chemistry'
    doc.core_properties.comments='AI-assisted manuscript draft; author review and approval are not asserted. Display equations are rendered images; source notation is supplied.'

def table(doc,lines,source):
    rows=[[x.strip() for x in l.strip().strip('|').split('|')] for l in lines]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',x) for x in r)]
    n=len(rows[0]);t=doc.add_table(rows=len(rows),cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if n==2:weights=[.75,.25]
    elif n==3:weights=[.44,.28,.28]
    else:weights=[.24,.16]+[.6/(n-2)]*(n-2)
    if n==3 and any('source' in c.lower() or '源码' in c for c in rows[0]):weights=[.28,.36,.36]
    for col,w in zip(t.columns,weights):col.width=Inches(6.7*w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for i,row in enumerate(rows):
        for j,value in enumerate(row):
            c=t.cell(i,j);c.width=Inches(6.7*weights[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');marg.append(e)
            cp.append(marg)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');cp.append(shade)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.08
            # Keep compact numerical tables intact; long ledgers repeat headers.
            if len(rows)<=10 and i<len(rows)-1:p.paragraph_format.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or '[' in value else WD_ALIGN_PARAGRAPH.CENTER
            inline(p,value,source,9.5)
            if i==0:
                for r in p.runs:r.bold=True
        rp=t.rows[i]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');rp.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader');rp.append(repeat)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(0);font(p.add_run(''),3)

def build(name):
    source=BASE/'rendered'/(name+'.md');text=source.read_text('utf-8')
    if '[@' in text or 'VERIFIED_REFERENCES' in text:raise ValueError('Unresolved references')
    text=DISPLAY.sub(lambda m:'\n@@EQ:'+equation_id((m.group(1) or m.group(2)).strip())+'\n',text)
    eq=json.loads((EQUATIONS/'equations_manifest.json').read_text('utf-8'))
    doc=Document();style_doc(doc,'chinese' in name,'supporting' in name)
    lines=text.splitlines();i=0;figures=0;equations=0;tables=0;is_refs=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line or line.startswith('<!--'):continue
        if line.startswith('```'):
            code=[]
            while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
            i+=1
            for raw in code:
                for wrapped in textwrap.wrap(raw,width=92,break_long_words=False,break_on_hyphens=False) or ['']:
                    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                    font(p.add_run(wrapped),8.5,mono=True)
            continue
        if line.startswith('|'):
            ls=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):ls.append(lines[i]);i+=1
            table(doc,ls,source);tables+=1;continue
        if line.startswith('@@EQ:'):
            key=line[5:];p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(EQUATIONS/(key+'.png')),width=Inches(eq[key]['width_inches']))
            p.paragraph_format.keep_together=True;equations+=1;continue
        im=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',line)
        if im:
            ip=(source.parent/im.group(2)).resolve();w,h=Image.open(ip).size
            width=min(6.7,4.1*w/h)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ip),width=Inches(width))
            p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4)
            figures+=1;continue
        h=re.match(r'^(#{1,4})\s+(.*)',line)
        if h:
            level=len(h.group(1));title=h.group(2)
            p=doc.add_paragraph(style='Title' if level==1 else 'Heading '+str(level-1));inline(p,title,source)
            if level==1:doc.core_properties.title=title
            is_refs=title in ['References','参考文献','S14. References','S14. 参考文献']
            continue
        paragraph=[line]
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||```|!\[|@@EQ:|[-*] |\d+\. )',lines[i]):
            paragraph.append(lines[i].strip());i+=1
        content=' '.join(paragraph)
        plain_caption=content.replace('**','')
        caption=bool(re.match(r'^(Figure|Table) S?\d+\.|^[图表]\s*S?\d+[.．。]',plain_caption))
        p=doc.add_paragraph(style='Caption' if caption else 'Normal')
        if caption:
            p.paragraph_format.keep_together=True
            if plain_caption.startswith(('Table ','表')):p.paragraph_format.keep_with_next=True
        if is_refs:
            p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)
            p.paragraph_format.line_spacing=1.08;p.paragraph_format.space_after=Pt(5)
        elif re.match(r'^[-*] ',content):content='• '+content[2:]
        inline(p,content,source,10 if is_refs else None)
    out=BASE/'documents';out.mkdir(exist_ok=True)
    target=out/(name+'.docx');doc.save(target)
    return {'name':name,'path':target.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'figures':figures,'equations':equations,'tables':tables,'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest()}

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--prepare-equations',action='store_true');ap.add_argument('--only',choices=NAMES,nargs='+');args=ap.parse_args()
    if args.prepare_equations:prepare()
    else:
        selected=args.only or NAMES
        record_path=BASE/'results'/'document_build.json'
        previous=json.loads(record_path.read_text('utf-8'))['documents'] if args.only and record_path.exists() else []
        by_name={d['name']:d for d in previous}
        by_name.update({name:build(name) for name in selected})
        result=[by_name[name] for name in NAMES if name in by_name]
        (BASE/'results'/'document_build.json').write_text(json.dumps({'documents':result,'equation_representation':'Matplotlib mathtext high-resolution images; editable LaTeX source supplied in assets/equations/equations.json'},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        print(json.dumps(result,ensure_ascii=False,indent=2))
